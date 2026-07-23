import json, re, sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
out = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    items = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        art, nr, keep = (row + (None, None, None))[:3]
        if art is None or str(art).strip() == '':
            continue
        items.append({
            'art_name': str(art).strip(),
            'nr_linii': nr,
            'pastreaza': str(keep or 'DA').strip().upper(),
        })
    key = 'website' if 'magazin' in sheet_name.lower() else 'marketplace'
    out[key] = items
wb.close()
sys.stdout.write(json.dumps(out, ensure_ascii=False))