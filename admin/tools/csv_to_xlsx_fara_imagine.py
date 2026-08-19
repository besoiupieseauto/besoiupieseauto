import csv
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

prod_csv, brand_csv, xlsx = sys.argv[1], sys.argv[2], sys.argv[3]

wb = Workbook(write_only=True)

ws_b = wb.create_sheet("Branduri")
ws_p = wb.create_sheet("Produse")

header_font = Font(bold=True)

with open(brand_csv, "r", encoding="utf-8-sig", newline="") as f:
    reader = csv.reader(f, delimiter=";")
    first = True
    for row in reader:
        if first:
            ws_b.append(row)
            first = False
        else:
            out = []
            for i, v in enumerate(row):
                if i > 0:
                    try:
                        out.append(int(v))
                    except ValueError:
                        out.append(v)
                else:
                    out.append(v)
            ws_b.append(out)

with open(prod_csv, "r", encoding="utf-8-sig", newline="") as f:
    reader = csv.reader(f, delimiter=";")
    first = True
    for row in reader:
        if first:
            ws_p.append(row)
            first = False
            continue
        if len(row) >= 7:
            try:
                row[6] = int(row[6])
            except ValueError:
                pass
        ws_p.append(row)

# write_only cannot set column widths easily after; still save
wb.save(xlsx)
print("OK", xlsx)
