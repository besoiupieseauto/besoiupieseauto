"""
Import Autonet QWP (CSV sau XLSX filtrat):
  - match ReferenceBrand+RefNr în besoiu_tecdoc_y1998
  - clone cartelă QWP+ArtNr în besoiu_tecdoc_base
  - UPSERT toate rândurile în shop.autonet_qwp_data
"""

from __future__ import annotations

import argparse
import csv
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from code_normalizer import normalize_code  # noqa: E402
from matcher import ProductMatcher  # noqa: E402
from paths import load_settings  # noqa: E402
from qwp_shop_map import get_qwp_shop_map  # noqa: E402
from qwp_tecdoc_writer import get_qwp_tecdoc_writer  # noqa: E402

DEFAULT_CSV = (
    Path(__file__).resolve().parents[3]
    / "Backend"
    / "storage"
    / "supplier_feeds"
    / "autonet"
    / "Autonet QWP-update filtrare.xlsx"
)

FALLBACK_DOWNLOADS = Path(r"c:\Users\Radu\Downloads\Autonet QWP-update filtrare.xlsx")


def _cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def iter_products(path: Path, limit: int | None = None, offset: int = 0):
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        yield from _iter_xlsx(path, limit=limit, offset=offset)
    else:
        yield from _iter_csv(path, limit=limit, offset=offset)


def _make_product(i: int, art: str, brand: str, ref: str, filename: str) -> dict:
    return {
        "row": i + 2,
        "sku_supplier": art,
        "art_nr": art,
        "ref_nr": ref,
        "brand": brand,
        "force_brand": "QWP",
        "supplier": "autonet",
        "supplier_profile": "autonet_qwp",
        "codes": [ref, normalize_code(ref)],
        "ean": "",
        "name": "",
        "match_code_only": True,
        "source_file": f"autonet/{filename}",
    }


def _iter_csv(path: Path, limit: int | None = None, offset: int = 0):
    with path.open(encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        for i, row in enumerate(reader):
            if i < offset:
                continue
            if limit is not None and (i - offset) >= limit:
                break
            art = _cell(row.get("ArtNr"))
            brand = _cell(row.get("ReferenceBrand"))
            ref = _cell(row.get("RefNr"))
            if not art or not ref:
                continue
            yield _make_product(i, art, brand, ref, path.name)


def _iter_xlsx(path: Path, limit: int | None = None, offset: int = 0):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Lipsește openpyxl — pip install openpyxl") from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [ _cell(h) for h in next(rows) ]
    # map case-insensitive
    idx = {h.upper(): n for n, h in enumerate(header)}
    need = ("ARTNR", "REFERENCEBRAND", "REFNR")
    for key in need:
        if key not in idx:
            wb.close()
            raise RuntimeError(f"Coloană lipsă în XLSX: {key}; header={header}")

    for i, row in enumerate(rows):
        if i < offset:
            continue
        if limit is not None and (i - offset) >= limit:
            break
        art = _cell(row[idx["ARTNR"]] if idx["ARTNR"] < len(row) else "")
        brand = _cell(row[idx["REFERENCEBRAND"]] if idx["REFERENCEBRAND"] < len(row) else "")
        ref = _cell(row[idx["REFNR"]] if idx["REFNR"] < len(row) else "")
        if not art or not ref:
            continue
        yield _make_product(i, art, brand, ref, path.name)
    wb.close()


def main() -> int:
    ap = argparse.ArgumentParser(description="Import QWP → TecDoc base + shop map")
    ap.add_argument("--csv", "--file", dest="file", type=Path, default=None)
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--offset", type=int, default=0)
    ap.add_argument("--batch", type=int, default=100)
    ap.add_argument("--skip-tecdoc-write", action="store_true")
    ap.add_argument("--skip-shop-map", action="store_true")
    args = ap.parse_args()

    path = args.file
    if path is None:
        path = DEFAULT_CSV if DEFAULT_CSV.is_file() else FALLBACK_DOWNLOADS
    if not path.is_file():
        print(f"Fișier lipsă: {path}")
        return 1

    settings = load_settings()
    print(
        f"FILE={path} offset={args.offset} limit={args.limit} "
        f"match_db={settings.get('tecdoc_mysql_db_qwp')} "
        f"write_db={settings.get('tecdoc_mysql_db_qwp_write')}"
    )

    matcher = ProductMatcher()
    writer = None if args.skip_tecdoc_write else get_qwp_tecdoc_writer()
    shop = None if args.skip_shop_map else get_qwp_shop_map()

    if writer and not writer.available:
        print("WARN: QwpTecdocWriter indisponibil")
    if shop and not shop.available:
        print("WARN: QwpShopMap indisponibil")

    totals = {
        "rows": 0,
        "exact": 0,
        "probable": 0,
        "no_match": 0,
        "written": 0,
        "written_no_compat": 0,
        "exists": 0,
        "map_upserted": 0,
        "error": 0,
    }
    batch: list[dict] = []
    t0 = time.time()

    def flush() -> None:
        nonlocal batch
        if not batch:
            return
        results = matcher.match_batch(batch)
        for r in results:
            st = r.get("status")
            if st in totals:
                totals[st] += 1
        if shop and shop.available:
            m = shop.upsert_from_match_inputs(results)
            totals["map_upserted"] += int(m.get("upserted") or 0)
        if writer and writer.available:
            w = writer.write_from_match_results(results)
            totals["written"] += int(w.get("written") or 0)
            totals["written_no_compat"] += int(w.get("written_no_compat") or 0)
            totals["exists"] += int(w.get("exists") or 0)
            totals["error"] += int(w.get("error") or 0)
        totals["rows"] += len(batch)
        elapsed = max(time.time() - t0, 0.001)
        print(
            f"... rows={totals['rows']} exact={totals['exact']} probable={totals['probable']} "
            f"no_match={totals['no_match']} written={totals['written']} "
            f"written_no_compat={totals['written_no_compat']} exists={totals['exists']} "
            f"map={totals['map_upserted']} rate={totals['rows']/elapsed:.1f}/s"
        )
        batch = []

    for product in iter_products(path, limit=args.limit, offset=args.offset):
        batch.append(product)
        if len(batch) >= max(1, args.batch):
            flush()
    flush()

    print("DONE", totals)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
